"""Restructures the first section of Radio_over_Fiber_Reviewed_Final_v3.pptx into 12 slides.

Keeps the original template (blank layout, white background, Arial, black text).
The three motivation slides are taken from Radio_over_Fiber_Reviewed_Final_v4.pptx,
which uses the identical template.

Usage: python3 restructure_v3.py v3.pptx v4.pptx out.pptx [figs_dir]
figs_dir: redrawn arch.png / block.png / rau.png from make_v3_figures.py (optional).
"""
import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

V3, V4, OUT = sys.argv[1:4]
FIGS = sys.argv[4] if len(sys.argv) > 4 else None
prs = Presentation(V3)
src = Presentation(V4)

BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x59, 0x59, 0x59)
FRAME = RGBColor(0xA6, 0xA6, 0xA6)
CALLOUT_FILL = RGBColor(0xF2, 0xF2, 0xF2)
FONT = "Arial"

orig = list(prs.slides)  # v3 slides, 0-based


def shape(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def set_runs(p, parts, size, color=BLACK):
    """parts: list of (text, bold, italic)."""
    for text, bold, italic in parts:
        r = p.add_run()
        r.text = text
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
        f.color.rgb = color


def fill_text(tf, paras, size, space_after=6, color=BLACK):
    """paras: list of part-lists (see set_runs). Replaces all existing text."""
    tf.clear()
    tf.word_wrap = True
    for i, parts in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        set_runs(p, parts, size, color)


def textbox(slide, x, y, w, h, paras, size, space_after=4, color=BLACK, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.05)
    tf.vertical_anchor = anchor
    fill_text(tf, paras, size, space_after, color)
    return tb


def frame(pic):
    pic.line.color.rgb = FRAME
    pic.line.width = Pt(0.75)


def card(slide, x, y, w, h, fill=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    style = s._element.find(qn("p:style"))  # drop theme style (shadow, white text)
    if style is not None:
        s._element.remove(style)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    s.line.color.rgb = FRAME
    s.line.width = Pt(0.75)
    return s


def replace_image(slide, name, png):
    """Swap the picture's image data in place, keeping its position, size and crop."""
    if not FIGS:
        return
    pic = shape(slide, name)
    rid = pic._element.blipFill.find(qn("a:blip")).get(qn("r:embed"))
    slide.part.related_part(rid)._blob = open(f"{FIGS}/{png}", "rb").read()


def set_title(slide, text):
    tf = shape(slide, "TextBox 1").text_frame
    tf.paragraphs[0].runs[0].text = text
    for r in tf.paragraphs[0].runs[1:]:
        r.text = ""


# ---------------------------------------------------------------- import v4 motivation slides
blank = orig[0].slide_layout


def import_slide(s_src):
    new = prs.slides.add_slide(blank)
    for ph in list(new.placeholders):
        ph._element.getparent().remove(ph._element)
    for el in s_src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:sp"), qn("p:graphicFrame"), qn("p:grpSp"), qn("p:cxnSp")):
            new.shapes._spTree.append(copy.deepcopy(el))
    return new


why = import_slide(src.slides[2])      # "Why Do We Need Radio over Fiber?"
coax = import_slide(src.slides[14])    # "RoF vs. Coaxial Cable and Microwave Links"
loss = import_slide(src.slides[15])    # "Example: Signal Loss over Distance"

# ---------------------------------------------------------------- 1. Title
s1 = orig[0]
tb3 = shape(s1, "TextBox 3")
tb3._element.getparent().remove(tb3._element)

# ---------------------------------------------------------------- Definition & principle (v3 slides 3 + 4)
d = orig[2]
set_title(d, "Radio over Fiber: Definition & Principle of Operation")
body = shape(d, "TextBox 2")
body.top, body.height = Inches(1.35), Inches(4.0)
fill_text(body.text_frame, [
    [("• Core concept: ", True, False),
     ("RoF transports an analog RF or IF waveform directly over an optical fiber carrier, rather than baseband digital bits.", False, False)],
    [("• Downlink (E/O → O/E): ", True, False),
     ("at the central station, the RF/IF signal modulates a CW laser carrier (direct modulation or an external MZM). "
      "At the remote unit, a photodiode recovers the RF/IF signal for filtering, amplification and antenna transmission.", False, False)],
    [("• Uplink: ", True, False),
     ("follows the reverse conversion path, from the remote antenna back to the central processing unit.", False, False)],
    [("• Multiplexing: ", True, False),
     ("wavelength-division multiplexing (WDM) lets several RF channels share one fiber on different wavelengths (λ1 … λn).", False, False)],
], 20, space_after=12)
c = card(d, 0.6, 5.55, 12.1, 0.95, CALLOUT_FILL)
tf = c.text_frame
tf.margin_left = tf.margin_right = Inches(0.2)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.word_wrap = True
fill_text(tf, [[("Critical distinction: ", True, False),
                ("RoF removes the transport bottleneck, but it does not create extra radio spectrum or wireless air-interface capacity by itself.", False, True)]],
          17, space_after=0)

# ---------------------------------------------------------------- Network architecture (v3 slide 5)
a = orig[4]
replace_image(a, "Picture 2", "arch.png")
pic = shape(a, "Picture 2")
frame(pic)
cap = shape(a, "TextBox 3")
cap.left, cap.width = pic.left, pic.width
cap.top = pic.top + pic.height + Inches(0.15)
cap.text_frame.margin_left = 0

# ---------------------------------------------------------------- Block diagram + building blocks (v3 slides 6 + 9)
b = orig[5]
replace_image(b, "Picture 2", "block.png")
pic = shape(b, "Picture 2")
ratio = pic.width / pic.height
pic.height = Inches(3.55)
pic.width = int(pic.height * ratio)
pic.left = int((prs.slide_width - pic.width) / 2)
pic.top = Inches(1.3)
frame(pic)
old = shape(b, "TextBox 3")
old._element.getparent().remove(old._element)
left = [
    [("CW laser: ", True, False), ("continuous-wave source that provides the optical carrier.", False, False)],
    [("MZM: ", True, False), ("Mach-Zehnder modulator; varies light intensity with the RF signal (alternative: direct laser modulation).", False, False)],
    [("WDM MUX / DEMUX: ", True, False), ("combine and separate channels λ1 … λn when they share one fiber.", False, False)],
]
right = [
    [("Photodiode: ", True, False), ("converts light back into the electrical RF signal at the base station (remote antenna site).", False, False)],
    [("Filter and PA: ", True, False), ("remove out-of-band products and amplify the RF signal before the antenna.", False, False)],
    [("Downlink shown; ", True, False), ("the uplink follows the reverse conversion path.", False, False)],
]
lx = (prs.slide_width - pic.width) / 2 / 914400
lw = pic.width / 914400
legend_top = 1.3 + 3.55 + 0.2
card(b, lx, legend_top, lw, 1.75)
textbox(b, lx + 0.12, legend_top + 0.08, lw / 2 - 0.2, 1.6, left, 13, space_after=5)
textbox(b, lx + lw / 2 + 0.08, legend_top + 0.08, lw / 2 - 0.2, 1.6, right, 13, space_after=5)

# ---------------------------------------------------------------- RAU types (v3 slide 7)
r = orig[6]
replace_image(r, "Picture 2", "rau.png")
frame(shape(r, "Picture 2"))
fill_text(shape(r, "TextBox 3").text_frame, [
    [("BBoF — baseband over fiber", True, False)],
    [("RAU: modulator + mixer + LO. Converts baseband up to RF.", False, False)],
    [("IFoF — intermediate frequency over fiber", True, False)],
    [("RAU: mixer + LO. Converts IF up to RF.", False, False)],
    [("RFoF — RF over fiber", True, False)],
    [("RAU: photodiode + TIA + PA. No frequency conversion; simplest RAU.", False, False)],
    [("All types", True, False)],
    [("Downlink: photodiode, TIA, PA.  Uplink: LNA, laser diode (LD).", False, False)],
], 18, space_after=4)
tf = shape(r, "TextBox 3").text_frame
for i, p in enumerate(tf.paragraphs):  # extra gap before each group heading
    if i and i % 2 == 0:
        p.space_before = Pt(14)

# ---------------------------------------------------------------- Analog vs digital table (v3 slide 8)
t = shape(orig[7], "Table 2").table
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def cell_border(cell, color="A6A6A6", w=9525):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for e in tcPr.findall(qn(tag)):
            tcPr.remove(e)
    # schema order: lnL, lnR, lnT, lnB must precede fill elements
    for i, tag in enumerate(("lnL", "lnR", "lnT", "lnB")):
        ln = etree.Element(qn("a:" + tag), w=str(w), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=color)
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
        tcPr.insert(i, ln)


for ri, row in enumerate(t.rows):
    for cell in row.cells:
        cell_border(cell)
        if ri == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xD9, 0xD9, 0xD9)

# ---------------------------------------------------------------- Advantages (v3 slides 10 + 11 + 12)
v = orig[9]
set_title(v, "Key Engineering Advantages of RoF")
for n in ("TextBox 2", "TextBox 3", "TextBox 4"):
    e = shape(v, n)
    e._element.getparent().remove(e._element)
items = [
    ("1. Ultra-low attenuation & extended reach",
     "≈ 0.2 dB/km at 1550 nm (ITU-T G.652 SMF), so one central site can serve a wide coverage area. Reach is bounded by noise, nonlinearity and dispersion."),
    ("2. Large transport bandwidth",
     "High fiber capacity, extended with WDM, supports multiple channels and operators on one fiber."),
    ("3. EMI immunity",
     "The optical fiber span is immune to electromagnetic interference. Over-the-air RF interference still applies."),
    ("4. Centralized processing, low field maintenance",
     "Complex processing stays at the BBU pool. Simpler remote units with fewer active parts reduce remote power and maintenance."),
    ("5. Protocol & modulation transparency",
     "A transparent optical link carries multiple radio generations and modulation formats without changing remote hardware."),
    ("6. Cost & installation efficiency",
     "Light, flexible cable simplifies installation in tunnels and buildings; existing dark fiber can be reused."),
]
cw, ch, gx, gy, x0, y0 = 3.9, 2.2, 0.2, 0.2, 0.6, 1.35
for i, (head, text) in enumerate(items):
    x = x0 + (i % 3) * (cw + gx)
    y = y0 + (i // 3) * (ch + gy)
    card(v, x, y, cw, ch)
    textbox(v, x + 0.12, y + 0.1, cw - 0.24, ch - 0.2, [[(head, True, False)], [(text, False, False)]], 15, space_after=6)
textbox(v, 0.6, 6.2, 12.1, 0.55, [[("Note: ", True, False),
        ("wireless capacity remains governed by radio spectrum, modulation efficiency and the number of radio units, not by fiber bandwidth alone.", False, False)]],
        14, space_after=0)
textbox(v, 0.6, 6.85, 12.1, 0.35, [[("Fiber attenuation: typical single-mode fiber at 1550 nm, ITU-T G.652 family.", False, False)]], 11, color=GREY)

# ---------------------------------------------------------------- order and delete
sld_lst = prs.slides._sldIdLst
ids = list(sld_lst)
new_ids = ids[len(orig):]  # why, coax, loss
keep = [ids[0], ids[1], new_ids[0], new_ids[1], new_ids[2],
        ids[2], ids[4], ids[5], ids[6], ids[7], ids[9], ids[12]]
for el in ids:
    if el not in keep:
        prs.part.drop_rel(el.get(qn("r:id")))
    sld_lst.remove(el)
for el in keep:
    sld_lst.append(el)

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
