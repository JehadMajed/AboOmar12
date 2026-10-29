"""Builds the full v3 deck: Part 1 (from restructure_v3.py) + Part 2 (DAS, deployment, limits) + backup.

Usage: python3 build_v3_full.py v3.pptx v4.pptx figs_v3 figs_part2 out.pptx
"""
import subprocess
import sys
import tempfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

V3, V4, FIGS1, FIGS2, OUT = sys.argv[1:6]
HERE = Path(__file__).resolve().parent

# ---------------------------------------------------------------- Part 1
part1 = Path(tempfile.mkdtemp()) / "part1.pptx"
subprocess.run([sys.executable, str(HERE / "restructure_v3.py"), V3, V4, str(part1), FIGS1], check=True)
prs = Presentation(str(part1))

# Part 2 starts with DAS, so the old DAS slide at the end of Part 1 is replaced by the new section.
sld_lst = prs.slides._sldIdLst
last = list(sld_lst)[-1]
prs.part.drop_rel(last.get(qn("r:id")))
sld_lst.remove(last)

INK = RGBColor(0x11, 0x11, 0x11)
GREY = RGBColor(0x59, 0x59, 0x59)
FRAME = RGBColor(0xA6, 0xA6, 0xA6)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
HEAD = RGBColor(0xD9, 0xD9, 0xD9)
ACCENT = RGBColor(0xD9, 0x53, 0x4F)
FONT = "Arial"
blank = prs.slides[0].slide_layout


# ---------------------------------------------------------------- helpers (same style as Part 1)
def runs(p, parts, size, color=INK):
    for item in parts:
        text, bold, italic = (item + (False, False))[:3] if isinstance(item, tuple) else (item, False, False)
        r = p.add_run()
        r.text = text
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
        f.color.rgb = color


def textbox(slide, x, y, w, h, paras, size, gap=6, color=INK, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.05)
    for i, parts in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        runs(p, parts, size, color)
    return tb


def box(slide, x, y, w, h, fill=None, line=FRAME):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    st = s._element.find(qn("p:style"))
    if st is not None:
        s._element.remove(st)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    s.line.color.rgb = line
    s.line.width = Pt(0.75)
    return s


def new_slide(title, notes=""):
    s = prs.slides.add_slide(blank)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    textbox(s, 0.6, 0.35, 12.1, 0.9, [[(title, True)]], 34)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def figure(slide, path, x, y, w, h, frame=True):
    """Place an image inside the box, keeping its aspect ratio; returns (x, y, w, h) actually used."""
    pw, ph = Image.open(path).size
    r = pw / ph
    fw, fh = w, w / r
    if fh > h:
        fh, fw = h, h * r
    fx = x + (w - fw) / 2
    pic = slide.shapes.add_picture(str(path), Inches(fx), Inches(y), Inches(fw), Inches(fh))
    if frame:
        pic.line.color.rgb = FRAME
        pic.line.width = Pt(0.75)
    return fx, y, fw, fh


def source(slide, text):
    textbox(slide, 0.6, 6.9, 12.1, 0.45, [[(text,)]], 11, color=GREY)


def table(slide, rows, x, y, w, col_w, size=14, row_h=0.5):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    st = tblPr.find(qn("a:tableStyleId"))
    if st is not None:
        st.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.cell(ri, ci)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            runs(p, [(val, ri == 0 or ci == 0)], size)
            c.fill.solid()
            c.fill.fore_color.rgb = HEAD if ri == 0 else RGBColor(0xFF, 0xFF, 0xFF)
            tcPr = c._tc.get_or_add_tcPr()
            for k, tag in enumerate(("lnL", "lnR", "lnT", "lnB")):
                ln = etree.Element(qn("a:" + tag), w="9525", cap="flat", cmpd="sng", algn="ctr")
                sf = etree.SubElement(ln, qn("a:solidFill"))
                etree.SubElement(sf, qn("a:srgbClr"), val="A6A6A6")
                etree.SubElement(ln, qn("a:prstDash"), val="solid")
                tcPr.insert(k, ln)
    return tbl


F1, F2 = Path(FIGS1), Path(FIGS2)

# ================================================================ Part 2: DAS + C-RAN (real figures, each with its source)
LI = "Li, Xie, El-Hajjar, Hanzo, \"Analogue Radio Over Fiber for Next-Generation RAN,\" arXiv:2111.13851 (2021)"
HM = "Hadi & Murtaza, \"Fibre Wireless Distributed Antenna Systems for 5G and 6G Services,\" Electronics 12(1), 64 (2023), CC BY 4.0"


def fig_slide(title, img, src, bullets=None, notes="", img_w=12.1, img_h=4.9, text_size=16):
    """Title + figure (+ optional bullets on the right) + source line under the figure."""
    s = new_slide(title, notes)
    if bullets:
        fx, fy, fw, fh = figure(s, F2 / img, 0.6, 1.35, img_w, img_h)
        textbox(s, fx + fw + 0.3, 1.4, 12.7 - (fx + fw + 0.3), 5.3, bullets, text_size, gap=10)
    else:
        figure(s, F2 / img, 0.6, 1.35, 12.1, img_h)
    source(s, "Source: " + src)
    return s


B = lambda h, t: [(h, True), (t, False)]

# 1. Transition
fig_slide("Part 2: Where Is RoF Used? DAS and C-RAN", "hadi2023_fig1_fiber_wireless_das.png", HM + ", Fig. 1.",
          [B("In buildings (DAS): ", "fiber carries the radio signal from one source to many indoor antennas."),
           B("In the mobile network (C-RAN): ", "fiber fronthaul links a central BBU pool to remote radio heads."),
           B("Same physics as Part 1: ", "loss, noise, dynamic range and transparency become design decisions.")],
          "Transition. The figure shows both uses of fiber-wireless transport in one picture: indoor DAS on the right, optical fronthaul from the BBU on the left.",
          img_w=6.6, img_h=5.3)

# 2. What is a DAS
s = new_slide("What Is a Distributed Antenna System (DAS)?",
              "For students: one signal source, many low-power antennas close to users. Left: a real ceiling DAS antenna in a shopping mall. Right: remote antenna units around one base station.")
fx, fy, fw, fh = figure(s, F2 / "silva2023_mall_das_antenna.jpg", 0.6, 1.35, 3.6, 4.6)
figure(s, F2 / "paredes2020_fig1_rau_network.png", 4.5, 1.35, 4.9, 4.6)
textbox(s, 9.7, 1.4, 3.0, 4.8, [
    [("1  Signal source", True)], [("operator base station",)],
    [("2  Distribution", True)], [("fiber or coax to remote units",)],
    [("3  Antennas", True)], [("many low-power antennas close to users",)],
], 15, gap=4)
source(s, "Left: ceiling DAS antenna, Silva et al., Sensors 23, 4579 (2023). Right: Paredes-Páliz et al., \"Radio over Fiber: An Alternative Broadband Network Technology for IoT,\" Electronics 9(11), 1785 (2020), Fig. 1, CC BY 4.0.")

# 3. Why fiber instead of coax (native chart from datasheet values)
s = new_slide("Why Fiber Instead of Coax in a DAS?",
              "Values computed from the LMR-400 datasheet attenuation formula (dB/100 ft = 0.12229·√f + 0.00026·f). Coax loss grows with frequency; fiber loss does not depend on the RF band.")
freqs = [700, 1800, 2600, 3500]
vals = [round((0.122290 * f ** 0.5 + 0.000260 * f) / 0.3048, 1) for f in freqs]
cd = CategoryChartData()
cd.categories = [f"{f} MHz" for f in freqs]
cd.add_series("Coax loss per 100 m (dB)", vals)
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.35), Inches(6.6), Inches(4.9), cd).chart
ch.has_legend = False
pl = ch.plots[0]
pl.has_data_labels = True
pl.data_labels.number_format = "0.0\" dB\""
pl.data_labels.number_format_is_linked = False
pl.data_labels.font.size = Pt(14)
pl.data_labels.font.bold = True
pl.series[0].format.fill.solid()
pl.series[0].format.fill.fore_color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
ch.value_axis.has_title = True
ch.value_axis.axis_title.text_frame.text = "Coax loss per 100 m (dB)"
ch.font.size = Pt(13)
ch.font.name = FONT
textbox(s, 7.5, 1.4, 5.2, 5.0, [
    B("Coax: ", "loss rises with frequency; ≈ 2.4× higher at 3.5 GHz than at 700 MHz."),
    B("Fiber: ", "≈ 0.025 dB per 100 m, the same for every band and operator."),
    B("Trade-off: ", "remote units need E/O, O/E, amplifiers and power."),
    B("Uplink: ", "N remote units on one port raise the noise floor by ≈ 10·log10(N) dB."),
], 16, gap=10)
source(s, "Source: Times Microwave LMR-400 datasheet (attenuation formula); G.652.D single-mode fiber, ≈ 0.25 dB/km at 1550 nm.")

# 4. D-RAN to C-RAN
fig_slide("From D-RAN to C-RAN", "li2021_fig1_ran_evolution.png", LI + ", Fig. 1.", None,
          "RAN evolution: the baseband unit moves away from the antenna site. 3G separates BBU and RRH; 4G pools BBUs centrally (C-RAN) with fiber fronthaul (CPRI); 5G splits the BBU into CU and DU with eCPRI.",
          img_h=3.4)
textbox(prs.slides[-1], 0.6, 5.0, 12.1, 1.6, [
    B("C-RAN: ", "baseband processing is pooled at a central site; remote radio heads stay at the antennas."),
    B("Benefits: ", "shared processing, easier coordination between cells, less equipment and power at each site."),
    B("Cost: ", "every remote radio head needs a high-capacity, low-latency fronthaul link."),
], 16, gap=6)

# 5. Fronthaul bottleneck
fig_slide("The Fronthaul Bottleneck", "li2021_fig3_bandwidth.jpg", LI + ", Fig. 3.",
          [B("Digital fronthaul ", "sends sampled I/Q bits; its rate grows with bandwidth and antenna count."),
           B("Example (CPRI): ", "20 MHz LTE, 1 antenna = 1.23 Gb/s; 100 MHz, 64 antennas ≈ 315 Gb/s."),
           B("Analog RoF ", "needs far less fiber bandwidth: up to 40-fold less in the modeled case.")],
          "The plot from Li et al. compares aggregated fiber bandwidth needed by A-RoF and D-RoF as the bandwidth per wavelength grows. CPRI numbers follow the CPRI V7.0 rate formula (15-bit I/Q, 16/15 control words, 8B/10B).",
          img_w=6.6, img_h=5.0)

# 6. Functional splits
fig_slide("Functional Splits: Where to Cut the Base Station", "quintuna2021_fig_functional_splits.png",
          "Quintuna Rodriguez, Guillemin, Ferrieux, Thomas, \"Cloud-RAN functional split for an efficient fronthaul network,\" arXiv:2101.04216 (2021).",
          [B("Option 8: ", "everything central; time-domain I/Q on the fiber (CPRI). Highest rate."),
           B("Option 7.2: ", "low-PHY at the radio unit; frequency-domain I/Q (eCPRI, O-RAN). Lower rate."),
           B("Option 2: ", "PDCP central; packets on the fiber. Lowest rate, relaxed latency."),
           B("Analog RoF: ", "goes beyond Option 8: the RF waveform itself is on the fiber.")],
          "The higher the split option number, the more processing stays central and the higher the fronthaul rate and latency requirements.",
          img_w=5.6, img_h=5.3, text_size=15)

# 7. A-RoF vs D-RoF in C-RAN
fig_slide("Analog vs. Digital Fronthaul in C-RAN", "li2021_fig2_arof_vs_drof.png", LI + ", Fig. 2.",
          [B("A-RoF: ", "RF on the fiber; RRH has photodiode, filter, amplifier."),
           B("D-RoF: ", "ADC at the CU and DAC at every RRH (red blocks)."),
           B("Result: ", "A-RoF is simpler and more power-efficient at the RRH; D-RoF is more robust.")],
          "Left: A-RoF, right: D-RoF. The red blocks are what digital fronthaul adds: data converters and higher-rate optics.",
          img_w=7.4, img_h=5.0)

# 8. Trade-off: signal quality
fig_slide("The Trade-off: Signal Quality (EVM)", "li2021_fig5_evm.jpg", LI + ", Fig. 5 (1310 nm system).", None,
          "EVM versus laser power. D-RoF (red / cyan) reaches about 3-4 % EVM; A-RoF (blue) stays near 8 % and degrades with more WDM channels. For 256-QAM the 3GPP limit is 3.5 % (TS 38.104), so in this modeled case A-RoF meets 64-QAM but not 256-QAM.",
          img_h=3.6)
textbox(prs.slides[-1], 0.6, 5.1, 12.1, 1.5, [
    B("A-RoF: ", "≈ 8 % EVM: meets the 64-QAM limit, not the 3.5 % limit for 256-QAM (3GPP TS 38.104)."),
    B("D-RoF: ", "meets the 256-QAM limit, at the cost of much higher fiber bandwidth and power."),
], 16, gap=6)

# 9. Applications
s = new_slide("RoF Application Areas", "One line per area.")
cards = [("In-building & venues", "Fiber-fed DAS in buildings, stadiums, airports, tunnels and metros."),
         ("5G fronthaul (C-RAN)", "Commercial fronthaul is mainly digital (CPRI / eCPRI); analog RoF is at research and trial stage."),
         ("Public safety", "In-building coverage for emergency-responder radio."),
         ("Defense & aerospace", "Antenna remoting, radar, satellite ground stations."),
         ("Research", "RF-over-fiber in the SKA-low radio-telescope receiver (Perini et al., 2022).")]
for k, (h, t) in enumerate(cards):
    col, row = k % 3, k // 3
    x = 0.6 + col * 4.1 if row == 0 else 2.65 + col * 4.1
    y = 1.45 + row * 2.45
    box(s, x, y, 3.9, 2.2)
    textbox(s, x + 0.12, y + 0.1, 3.66, 2.0, [[(h, True)], [(t,)]], 16, gap=8)
source(s, "Sources: ITU-T G.9803 Amd. 2 (2022); Perini et al., J. Astron. Telesc. Instrum. Syst. 8(1), 011016 (2022).")

# 10. Limitations
s = new_slide("Limitations & Technical Challenges", "Practical limits first, then technical limits, each with its mitigation.")
table(s, [
    ["Challenge", "Impact", "Mitigation"],
    ["Cost & fiber availability", "New fiber and optoelectronics raise CAPEX", "Reuse existing fiber; share among operators"],
    ["Remote power", "Every remote unit needs power and backup", "Hybrid fiber–copper cable; local UPS"],
    ["Fronthaul capacity (C-RAN)", "Digital fronthaul rate grows with antennas and bandwidth", "Split 7.2x (eCPRI); analog RoF"],
    ["Latency", "Fronthaul distance limited by HARQ timing", "Keep BBU pool within the latency budget"],
    ["Nonlinearity & noise", "Limit dynamic range and EVM", "Optimized drive level; predistortion"],
    ["Chromatic dispersion", "RF power fading at high RF and long fiber", "1310 nm; single-sideband; IFoF"],
    ["Interoperability", "Analog interfaces are vendor-specific", "ITU-T G.9803 requirements; eCPRI / O-RAN"],
], 0.6, 1.35, 12.1, [3.2, 4.6, 4.3], size=14, row_h=0.62)

# 11. Summary
s = new_slide("Summary", "Five take-aways in the order of the talk.")
textbox(s, 0.6, 1.4, 12.1, 5.3, [
    [("• RoF carries the analog radio waveform over fiber: E/O at one end, O/E at the other.",)],
    [("• Fiber loss is low and band-independent; the link still adds conversion loss, noise and distortion.",)],
    [("• In buildings, a fiber-fed DAS uses RoF to feed many antennas from one signal source.",)],
    [("• In C-RAN, analog RoF cuts fronthaul bandwidth sharply but loses signal quality versus digital fronthaul.",)],
    [("• RoF moves the antenna closer to users; it does not add spectrum or air-interface capacity.", True)],
], 20, gap=14)

# 12. References
s = new_slide("References")
refs = [
    "ITU-T G.9803 (2018), Radio over fibre systems; Amendments 1 (2019) and 2 (2022). ITU-T G.Suppl.55 (10/2025). ITU-T G.652.",
    "Y. Li, Q. Xie, M. El-Hajjar, L. Hanzo, \"Analogue Radio Over Fiber for Next-Generation RAN: Challenges and Opportunities,\" arXiv:2111.13851, 2021.",
    "M. U. Hadi, G. Murtaza, \"Fibre Wireless Distributed Antenna Systems for 5G and 6G Services,\" Electronics 12(1), 64, 2023. DOI: 10.3390/electronics12010064.",
    "D. F. Paredes-Páliz et al., \"Radio over Fiber: An Alternative Broadband Network Technology for IoT,\" Electronics 9(11), 1785, 2020. DOI: 10.3390/electronics9111785.",
    "V. Quintuna Rodriguez et al., \"Cloud-RAN functional split for an efficient fronthaul network,\" arXiv:2101.04216, 2021.",
    "J. da L. A. Silva et al., \"Human Exposure to Non-Ionizing Radiation from Indoor DAS: Shopping Mall Measurement Analysis,\" Sensors 23, 4579, 2023.",
    "CPRI Specification V7.0 (2015); eCPRI Specification V2.0 (2019); 3GPP TS 38.104.",
    "Times Microwave Systems, LMR-400 coaxial cable datasheet.",
    "F. Perini et al., \"Radio frequency over fiber technology for SKA-low receiver,\" JATIS 8(1), 011016, 2022.",
    "SolFinder Research, \"Radio Over Fiber,\" video presentation (source of original material).",
]
textbox(s, 0.6, 1.35, 12.1, 5.7, [[(f"{k + 1}. {r}",)] for k, r in enumerate(refs)], 13, gap=6)

# Backup
fig_slide("Backup: 5G + Candidate 6G Demo on One Fiber", "hadi2023_fig2_5g_6g_demo.png", HM + ", Fig. 2.", None,
          "Laboratory demonstration only: a 5G-NR signal (IF1 = 1.14 GHz) and a candidate 6G THz signal (IF2 = 12 GHz) carried on one fiber.",
          img_h=5.2)
prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
